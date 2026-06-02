import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import glob, os

# Buscar el archivo original
candidatos = glob.glob(r'C:\Users\Mezcua\Downloads\Cat*.xlsx')
if not candidatos:
    print("No se encontro el archivo Excel en Descargas")
    exit(1)

src = candidatos[0]
print(f"Leyendo: {src}")

wb_src = openpyxl.load_workbook(src, data_only=True)
ws_src = wb_src.active
rows = list(ws_src.iter_rows(values_only=True))[1:]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Catalogo"

headers = ["tipo", "descripcion", "precio_unitario", "unidad", "referencia"]
fill = PatternFill("solid", fgColor="1E40AF")
font = Font(bold=True, color="FFFFFF")
for i, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=i, value=h)
    c.fill = fill
    c.font = font
    c.alignment = Alignment(horizontal="center")

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 55
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 10
ws.column_dimensions["E"].width = 25

n = 0
for row in rows:
    categoria  = str(row[1]).strip() if row[1] else ""
    elemento   = str(row[2]).strip() if row[2] else ""
    marca      = str(row[3]).strip() if row[3] else ""
    precio     = row[4]
    unidad_raw = str(row[5]).strip() if row[5] else "ud"

    if not elemento or "ategor" in categoria:
        continue

    descripcion = (elemento + " - " + marca) if marca else elemento

    unidad = "ud"
    if unidad_raw.lower() in ("metro", "m", "metros"):
        unidad = "m"
    elif unidad_raw.lower() in ("rollo", "rollos"):
        unidad = "rollo"

    try:
        precio_f = float(precio) if precio else 0.0
    except Exception:
        precio_f = 0.0

    ws.append(["Material", descripcion, precio_f, unidad, categoria])
    n += 1

out = r"C:\Users\Mezcua\Downloads\catalogo_railway.xlsx"
wb.save(out)
print(f"Guardado: {n} productos en {out}")
