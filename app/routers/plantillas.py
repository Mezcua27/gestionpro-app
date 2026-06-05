import secrets
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
from datetime import datetime
from io import BytesIO

from app.database import get_db
from app import models, schemas
from app.auth import get_current_user, get_effective_empresa_id

router = APIRouter()


def _get_plantilla(id: int, eid: int, db: Session) -> models.PlantillaPresupuesto:
    p = db.query(models.PlantillaPresupuesto).filter(
        models.PlantillaPresupuesto.id == id,
        models.PlantillaPresupuesto.empresa_id == eid
    ).first()
    if not p:
        raise HTTPException(status_code=404, detail="Plantilla no encontrada")
    return p


# ── CRUD plantillas ───────────────────────────────────────────────────────────

@router.get("/", response_model=List[schemas.PlantillaResponse])
def listar(request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    return db.query(models.PlantillaPresupuesto).filter(
        models.PlantillaPresupuesto.empresa_id == eid
    ).order_by(models.PlantillaPresupuesto.nombre).all()


@router.get("/{id}", response_model=schemas.PlantillaResponse)
def obtener(id: int, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    return _get_plantilla(id, eid, db)


@router.post("/", response_model=schemas.PlantillaResponse)
def crear(data: schemas.PlantillaCreate, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    p = models.PlantillaPresupuesto(**data.dict(), empresa_id=eid)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


@router.put("/{id}", response_model=schemas.PlantillaResponse)
def actualizar(id: int, data: schemas.PlantillaUpdate, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    p = _get_plantilla(id, eid, db)
    for k, v in data.dict(exclude_none=True).items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return p


@router.delete("/{id}")
def eliminar(id: int, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    p = _get_plantilla(id, eid, db)
    db.delete(p)
    db.commit()
    return {"ok": True}


# ── Líneas de plantilla ───────────────────────────────────────────────────────

@router.post("/{id}/lineas", response_model=schemas.LineaPlantillaResponse)
def añadir_linea(id: int, data: schemas.LineaPlantillaCreate, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    _get_plantilla(id, eid, db)
    linea = models.LineaPlantilla(**data.dict(), plantilla_id=id)
    db.add(linea)
    db.commit()
    db.refresh(linea)
    return linea


@router.post("/{id}/lineas/bulk")
def añadir_lineas_bulk(id: int, lineas: List[schemas.LineaPlantillaCreate], request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    p = _get_plantilla(id, eid, db)
    orden_base = len(p.lineas)
    for i, l in enumerate(lineas):
        datos = l.dict()
        datos['orden'] = orden_base + i
        db.add(models.LineaPlantilla(**datos, plantilla_id=id))
    db.commit()
    db.refresh(p)
    return {"ok": True, "añadidas": len(lineas)}


@router.put("/{id}/lineas/{linea_id}", response_model=schemas.LineaPlantillaResponse)
def editar_linea(id: int, linea_id: int, data: schemas.LineaPlantillaUpdate,
                 request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    _get_plantilla(id, eid, db)
    linea = db.query(models.LineaPlantilla).filter(
        models.LineaPlantilla.id == linea_id,
        models.LineaPlantilla.plantilla_id == id
    ).first()
    if not linea:
        raise HTTPException(status_code=404, detail="Línea no encontrada")
    for k, v in data.dict(exclude_none=True).items():
        setattr(linea, k, v)
    db.commit()
    db.refresh(linea)
    return linea


@router.delete("/{id}/lineas/{linea_id}")
def eliminar_linea(id: int, linea_id: int, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    _get_plantilla(id, eid, db)
    linea = db.query(models.LineaPlantilla).filter(
        models.LineaPlantilla.id == linea_id,
        models.LineaPlantilla.plantilla_id == id
    ).first()
    if not linea:
        raise HTTPException(status_code=404, detail="Línea no encontrada")
    db.delete(linea)
    db.commit()
    return {"ok": True}


# ── Usar plantilla → crear presupuesto ───────────────────────────────────────

@router.post("/{id}/usar", response_model=schemas.PresupuestoResponse)
def usar_plantilla(id: int, data: dict, request: Request, db: Session = Depends(get_db)):
    """Crea un nuevo presupuesto en borrador a partir de esta plantilla."""
    user = get_current_user(request, db)
    eid = get_effective_empresa_id(user, request)
    plantilla = _get_plantilla(id, eid, db)

    cliente_id = data.get("cliente_id")
    if cliente_id:
        c = db.query(models.Cliente).filter(
            models.Cliente.id == cliente_id,
            models.Cliente.empresa_id == eid
        ).first()
        if not c:
            raise HTTPException(status_code=400, detail="Cliente no válido")

    # Generar número correlativo
    year = datetime.utcnow().year
    count = db.query(models.Presupuesto).filter(
        models.Presupuesto.empresa_id == eid
    ).count() + 1
    numero = f"PRES-{year}-{count:04d}"

    titulo = data.get("titulo") or plantilla.nombre
    pres = models.Presupuesto(
        empresa_id=eid,
        cliente_id=cliente_id,
        numero=numero,
        titulo=titulo,
        estado="borrador",
        validez_dias=plantilla.validez_dias,
        notas=plantilla.notas,
        descuento=plantilla.descuento,
        iva=plantilla.iva,
        creado_por=user.id,
        token_cliente=secrets.token_urlsafe(32),
    )
    db.add(pres)
    db.flush()

    for i, l in enumerate(plantilla.lineas):
        db.add(models.LineaPresupuesto(
            presupuesto_id=pres.id,
            tipo=l.tipo,
            descripcion=l.descripcion,
            cantidad=l.cantidad,
            precio_unitario=l.precio_unitario,
            unidad=l.unidad,
            referencia=l.referencia,
            orden=i,
        ))

    db.commit()
    db.refresh(pres)
    return pres


# ── Importar plantilla desde Excel ───────────────────────────────────────────

@router.post("/importar-excel")
async def importar_excel(request: Request, archivo: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Importa una o varias plantillas desde Excel.
    Formato del fichero:
      Fila 1  : cabecera decorativa (se ignora)
      Filas 2-7: NOMBRE/CATEGORIA/NOTAS/IVA/DESCUENTO/VALIDEZ_DIAS  |  valor
      Fila 8  : vacía
      Fila 9  : cabecera de líneas (tipo|descripcion|cantidad|precio_unitario|unidad|referencia)
      Fila 10+: líneas de la plantilla
    """
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado.")

    user = get_current_user(request, db)
    eid  = get_effective_empresa_id(user, request)

    if not archivo.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx")

    contenido = await archivo.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo abrir el Excel: {e}")

    importadas = 0
    errores    = []

    for sheet_name in wb.sheetnames:
        ws   = wb[sheet_name]
        rows = [r for r in ws.iter_rows(values_only=True)]
        if len(rows) < 10:
            errores.append(f"Hoja '{sheet_name}': pocas filas, se omite.")
            continue

        # Leer metadatos (filas 2-7, col A=clave, col B=valor)
        meta = {}
        for row in rows[1:7]:
            if row[0]:
                meta[str(row[0]).strip().upper()] = str(row[1]).strip() if row[1] is not None else ""

        nombre = meta.get("NOMBRE", "").strip()
        if not nombre:
            errores.append(f"Hoja '{sheet_name}': sin NOMBRE, se omite.")
            continue

        try:
            iva          = float(meta.get("IVA", "21"))
            descuento    = float(meta.get("DESCUENTO", "0"))
            validez_dias = int(meta.get("VALIDEZ_DIAS", "30"))
        except ValueError:
            iva, descuento, validez_dias = 21.0, 0.0, 30

        # Detectar fila de cabecera de líneas buscando "tipo" en alguna columna
        header_row_idx = None
        for i, row in enumerate(rows):
            vals = [str(c).strip().lower() if c else "" for c in row]
            if "tipo" in vals and "descripcion" in vals:
                header_row_idx = i
                break

        if header_row_idx is None:
            errores.append(f"Hoja '{sheet_name}': no se encontró cabecera de líneas.")
            continue

        header = [str(c).strip().lower() if c else "" for c in rows[header_row_idx]]

        def col(name):
            try: return header.index(name)
            except ValueError: return None

        ci_tipo  = col("tipo")
        ci_desc  = col("descripcion")
        ci_cant  = col("cantidad")
        ci_prec  = col("precio_unitario")
        ci_uni   = col("unidad")
        ci_ref   = col("referencia")

        if ci_tipo is None or ci_desc is None or ci_prec is None:
            errores.append(f"Hoja '{sheet_name}': faltan columnas obligatorias.")
            continue

        # Crear la plantilla
        pt = models.PlantillaPresupuesto(
            empresa_id   = eid,
            nombre       = nombre,
            categoria    = meta.get("CATEGORIA") or None,
            notas        = meta.get("NOTAS") or None,
            iva          = iva,
            descuento    = descuento,
            validez_dias = validez_dias,
        )
        db.add(pt)
        db.flush()

        tipos_validos = {"Material", "Mano de Obra", "Varios"}
        orden = 0
        for row in rows[header_row_idx + 1:]:
            if not any(row):
                continue
            tipo = str(row[ci_tipo]).strip() if ci_tipo is not None and row[ci_tipo] else ""
            desc = str(row[ci_desc]).strip() if ci_desc is not None and row[ci_desc] else ""
            if not tipo or not desc:
                continue
            if tipo not in tipos_validos:
                tipo = "Material"
            try:
                precio = float(row[ci_prec]) if ci_prec is not None and row[ci_prec] is not None else 0.0
            except (ValueError, TypeError):
                precio = 0.0
            try:
                cantidad = float(row[ci_cant]) if ci_cant is not None and row[ci_cant] is not None else 1.0
            except (ValueError, TypeError):
                cantidad = 1.0
            unidad = str(row[ci_uni]).strip()  if ci_uni  is not None and row[ci_uni]  else "ud"
            ref    = str(row[ci_ref]).strip()  if ci_ref  is not None and row[ci_ref]  else None

            db.add(models.LineaPlantilla(
                plantilla_id    = pt.id,
                tipo            = tipo,
                descripcion     = desc,
                cantidad        = cantidad,
                precio_unitario = precio,
                unidad          = unidad,
                referencia      = ref or None,
                orden           = orden,
            ))
            orden += 1

        db.commit()
        importadas += 1

    return {"ok": True, "importadas": importadas, "errores": errores}


@router.get("/plantilla-excel")
def descargar_plantilla_ejemplo():
    """Descarga un Excel de ejemplo con el formato correcto para importar plantillas."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla"

    fill_h = PatternFill("solid", fgColor="1E40AF")
    font_h = Font(bold=True, color="FFFFFF", size=11)
    font_b = Font(bold=True, size=10)
    al_c   = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "PLANTILLA DE PRESUPUESTO - IMPORTACION"
    c.font  = Font(bold=True, color="FFFFFF", size=12)
    c.fill  = fill_h
    c.alignment = al_c

    meta = [
        ("NOMBRE",       "Cuadro Electrico Electrificacion Basica - Gama Alta"),
        ("CATEGORIA",    "Electricidad"),
        ("NOTAS",        "Instalacion electrica vivienda. Precios sin IVA."),
        ("IVA",          "21"),
        ("DESCUENTO",    "0"),
        ("VALIDEZ_DIAS", "30"),
    ]
    for i, (k, v) in enumerate(meta, 2):
        ws.cell(row=i, column=1, value=k).font = font_b
        ws.cell(row=i, column=2, value=v)
        ws.merge_cells(f"B{i}:F{i}")

    ws.append([])

    heads = ["tipo", "descripcion", "cantidad", "precio_unitario", "unidad", "referencia"]
    for col, h in enumerate(heads, 1):
        c = ws.cell(row=9, column=col, value=h)
        c.fill = fill_h; c.font = font_h; c.alignment = al_c

    ejemplos = [
        ("Material",    "Interruptor General Automatico 2P 40A", 1, 34.71, "ud", "A9F74240"),
        ("Material",    "Interruptor Diferencial 2P 40A 30mA Tipo AC", 1, 56.20, "ud", "A9Z05240"),
        ("Mano de Obra","Cuadro electrico completo instalado", 1, 330.58, "ud", ""),
        ("Varios",      "Certificado instalacion electrica CIE", 1, 206.61, "ud", ""),
    ]
    for e in ejemplos:
        ws.append(list(e))

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 20

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_ejemplo.xlsx"}
    )
