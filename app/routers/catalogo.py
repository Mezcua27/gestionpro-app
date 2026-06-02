import os
import uuid
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
from app.database import get_db
from app import models, schemas
from app.auth import get_current_user

router = APIRouter()

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static", "uploads", "catalogo")
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/gif"}
MAX_IMAGE_SIZE = 5 * 1024 * 1024  # 5 MB


@router.get("/", response_model=List[schemas.CatalogoItemResponse])
def listar(request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    return db.query(models.CatalogoItem).filter(
        models.CatalogoItem.empresa_id == user.empresa_id,
        models.CatalogoItem.activo == True
    ).order_by(models.CatalogoItem.tipo, models.CatalogoItem.descripcion).all()


@router.post("/", response_model=schemas.CatalogoItemResponse)
def crear(data: schemas.CatalogoItemCreate, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    item = models.CatalogoItem(**data.dict(), empresa_id=user.empresa_id)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@router.put("/{id}", response_model=schemas.CatalogoItemResponse)
def actualizar(id: int, data: schemas.CatalogoItemCreate, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    item = db.query(models.CatalogoItem).filter(
        models.CatalogoItem.id == id,
        models.CatalogoItem.empresa_id == user.empresa_id
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Ítem no encontrado")
    for k, v in data.dict().items():
        setattr(item, k, v)
    db.commit()
    db.refresh(item)
    return item


@router.delete("/{id}")
def eliminar(id: int, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    item = db.query(models.CatalogoItem).filter(
        models.CatalogoItem.id == id,
        models.CatalogoItem.empresa_id == user.empresa_id
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Ítem no encontrado")
    if item.foto:
        foto_path = os.path.join(UPLOAD_DIR, item.foto)
        if os.path.exists(foto_path):
            os.remove(foto_path)
    item.activo = False
    db.commit()
    return {"ok": True}


# ── Foto ────────────────────────────────────────────────────────────────────

@router.post("/{id}/foto")
async def subir_foto(id: int, request: Request, foto: UploadFile = File(...), db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    item = db.query(models.CatalogoItem).filter(
        models.CatalogoItem.id == id,
        models.CatalogoItem.empresa_id == user.empresa_id
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Ítem no encontrado")

    if foto.content_type not in ALLOWED_IMAGE_TYPES:
        raise HTTPException(status_code=400, detail="Formato no permitido. Usa JPG, PNG o WebP.")

    contenido = await foto.read()
    if len(contenido) > MAX_IMAGE_SIZE:
        raise HTTPException(status_code=400, detail="La imagen no puede superar 5 MB.")

    ext = foto.filename.rsplit(".", 1)[-1].lower() if "." in foto.filename else "jpg"
    nombre_archivo = f"{user.empresa_id}_{id}_{uuid.uuid4().hex[:8]}.{ext}"

    # Borrar foto anterior si existe
    if item.foto:
        anterior = os.path.join(UPLOAD_DIR, item.foto)
        if os.path.exists(anterior):
            os.remove(anterior)

    ruta = os.path.join(UPLOAD_DIR, nombre_archivo)
    with open(ruta, "wb") as f:
        f.write(contenido)

    item.foto = nombre_archivo
    db.commit()
    return {"ok": True, "foto": f"/static/uploads/catalogo/{nombre_archivo}"}


@router.delete("/{id}/foto")
def eliminar_foto(id: int, request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    item = db.query(models.CatalogoItem).filter(
        models.CatalogoItem.id == id,
        models.CatalogoItem.empresa_id == user.empresa_id
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Ítem no encontrado")
    if item.foto:
        ruta = os.path.join(UPLOAD_DIR, item.foto)
        if os.path.exists(ruta):
            os.remove(ruta)
        item.foto = None
        db.commit()
    return {"ok": True}


# ── Excel ───────────────────────────────────────────────────────────────────

@router.get("/plantilla-excel")
def descargar_plantilla():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from io import BytesIO
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catálogo"

    headers = ["tipo", "descripcion", "precio_unitario", "unidad", "referencia"]
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20

    # Fila de ejemplo
    ws.append(["Material", "Cable eléctrico 2.5mm²", 1.20, "m", "CAB-2.5"])
    ws.append(["Mano de Obra", "Oficial electricista", 35.00, "h", ""])
    ws.append(["Varios", "Transporte", 50.00, "ud", "TRANS-01"])

    # Hoja de ayuda
    ws2 = wb.create_sheet("Ayuda")
    ws2["A1"] = "Valores válidos para la columna 'tipo':"
    ws2["A2"] = "Material"
    ws2["A3"] = "Mano de Obra"
    ws2["A4"] = "Varios"
    ws2["A6"] = "Notas:"
    ws2["A7"] = "- precio_unitario: usa punto (.) como separador decimal"
    ws2["A8"] = "- referencia: opcional, puede dejarse vacío"
    ws2["A9"] = "- Las fotos se añaden desde la aplicación una vez importado el catálogo"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_catalogo.xlsx"}
    )


@router.post("/importar-excel")
async def importar_excel(request: Request, archivo: UploadFile = File(...), db: Session = Depends(get_db)):
    try:
        import openpyxl
        from io import BytesIO
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado en el servidor.")

    user = get_current_user(request, db)

    if not archivo.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .xlsx")

    try:
        contenido = await archivo.read()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error leyendo el archivo: {e}")

    try:
        wb = openpyxl.load_workbook(BytesIO(contenido), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo abrir el Excel: {e}")

    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error leyendo las filas: {e}")

    if not rows:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")

    # Detectar cabecera — admite con o sin acentos
    header_row = [str(c).strip().lower() if c else "" for c in rows[0]]

    # Normalizar nombres de columna (quitar acentos para mayor compatibilidad)
    def norm(s):
        import unicodedata
        return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')

    header_norm = [norm(h) for h in header_row]

    required = {"tipo", "descripcion", "precio_unitario"}
    if not required.issubset(set(header_norm)):
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas. Necesarias: {', '.join(required)}. "
                   f"Columnas encontradas: {', '.join(h for h in header_norm if h)}. "
                   f"Descarga la plantilla para ver el formato correcto."
        )

    col = {name: i for i, name in enumerate(header_norm)}
    tipos_validos = {"Material", "Mano de Obra", "Varios"}
    importados = 0
    errores = []

    try:
        for fila_num, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            tipo_raw = row[col["tipo"]]
            desc_raw = row[col["descripcion"]]
            tipo = str(tipo_raw).strip() if tipo_raw else ""
            descripcion = str(desc_raw).strip() if desc_raw else ""

            if tipo not in tipos_validos:
                errores.append(f"Fila {fila_num}: tipo '{tipo}' no válido.")
                continue
            if not descripcion:
                errores.append(f"Fila {fila_num}: descripción vacía.")
                continue

            try:
                precio_raw = row[col["precio_unitario"]]
                precio = float(precio_raw) if precio_raw is not None else 0.0
            except (ValueError, TypeError):
                precio = 0.0

            unidad = "ud"
            if "unidad" in col:
                val = row[col["unidad"]]
                if val:
                    unidad = str(val).strip() or "ud"

            referencia = None
            if "referencia" in col:
                val = row[col["referencia"]]
                if val and str(val).strip() not in ("", "None"):
                    referencia = str(val).strip()

            item = models.CatalogoItem(
                empresa_id=user.empresa_id,
                tipo=tipo,
                descripcion=descripcion,
                precio_unitario=precio,
                unidad=unidad,
                referencia=referencia,
            )
            db.add(item)
            importados += 1

        if importados:
            db.commit()

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al guardar en base de datos: {str(e)}")

    return {
        "ok": True,
        "importados": importados,
        "errores": errores,
        "mensaje": f"Se importaron {importados} ítems." + (f" {len(errores)} filas omitidas." if errores else "")
    }
